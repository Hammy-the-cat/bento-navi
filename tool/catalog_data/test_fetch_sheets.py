import importlib.util
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch
import openpyxl

spec = importlib.util.spec_from_file_location('fetch_catalog_sheets', Path(__file__).with_name('fetch_sheets.py'))
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class FetchTests(unittest.TestCase):
    def test_full_export_expands_rows_and_preserves_phone_display(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = '宮崎県'
        sheet.append(['店舗ID', '店舗名', '電話番号'])
        sheet.append(['S1', '弁当屋', 985123456])
        sheet['C2'].number_format = '0000-00-0000'
        sheet.cell(100, 1, 'S100')
        data = BytesIO()
        workbook.save(data)
        with patch.object(module, 'urlopen', return_value=BytesIO(data.getvalue())):
            tabs = module.read_book({'id':'book', 'title':'本'}, [{'spreadsheetId':'book', 'sheetId':0, 'title':'宮崎県', 'book':'本'}])
        self.assertEqual(len(tabs[0]['values']), 100)
        self.assertEqual(tabs[0]['values'][1][2], '0985-12-3456')
        self.assertTrue(tabs[0]['range'].endswith(':C100'))

    def test_missing_registered_tab_and_unregistered_master_fail(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = '新地域'
        workbook.active.append(['店舗ID','店舗名'])
        data = BytesIO()
        workbook.save(data)
        with patch.object(module, 'urlopen', return_value=BytesIO(data.getvalue())):
            with self.assertRaisesRegex(ValueError, 'Missing registered'):
                module.read_book({'id':'book','title':'本'}, [{'title':'宮崎県'}])
        with patch.object(module, 'urlopen', return_value=BytesIO(data.getvalue())):
            with self.assertRaisesRegex(ValueError, 'New shop tab'):
                module.read_book({'id':'book','title':'本'}, [])
