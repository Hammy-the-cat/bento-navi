"""Download complete workbooks using their existing read access; never change sharing."""
import argparse
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, date, timezone
from io import BytesIO
import json
from pathlib import Path
import time
from urllib.request import urlopen, Request
from zipfile import ZipFile

import openpyxl
from openpyxl.utils import get_column_letter


def cell_text(cell):
    value = cell.value
    if value is None:
        return ''
    if cell.data_type == 'e':
        raise ValueError('Spreadsheet formula error')
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, (int, float)):
        # Preserve leading zeros in explicitly formatted numeric phone/ID cells.
        fmt = cell.number_format
        if fmt and set(fmt) <= {'0', '-'} and '0' in fmt and float(value).is_integer():
            digits = str(int(value)).zfill(fmt.count('0'))
            chars = iter(digits)
            return ''.join(next(chars) if c == '0' else c for c in fmt)
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value)


def read_book(book, tabs):
    url = f"https://docs.google.com/spreadsheets/d/{book['id']}/export?format=xlsx"
    for attempt in range(3):
        try:
            with urlopen(Request(url, headers={'User-Agent': 'BentoNavi-SheetSync/1.0'}), timeout=90) as response:
                raw = response.read(20_000_001)
            if len(raw) > 20_000_000 or not raw.startswith(b'PK'):
                raise ValueError('Invalid workbook export (check source read access)')
            with ZipFile(BytesIO(raw)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 100_000_000:
                    raise ValueError('Workbook too large')
            workbook = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True)
            break
        except Exception:
            if attempt == 2:
                raise
            time.sleep(2 ** attempt)
    result = []
    expected = {t['title']: t for t in tabs}
    if not set(expected) <= set(workbook.sheetnames):
        raise ValueError(f"Missing registered tabs in {book['title']}")
    for sheet in workbook:
        sheet.calculate_dimension(force=True)
        if sheet.max_row > 50000 or sheet.max_column > 100:
            raise ValueError('Sheet dimensions exceed catalog limits')
        headers = [cell_text(c) for c in next(sheet.iter_rows(max_row=1))]
        if sheet.title not in expected:
            # Supporting evidence/call lists repeat IDs but are not shop masters.
            support_tabs = {'調査メモ', '裏付け一覧', '網羅状況', 'スポーツ施設別', '電話確認リスト'}
            if sheet.title not in support_tabs and '店舗ID' in headers and '店舗名' in headers:
                raise ValueError(f'New shop tab must be registered: {book["title"]}/{sheet.title}')
            continue
        values = [[cell_text(c) for c in row] for row in sheet.iter_rows()]
        # XLSX export includes the complete sheet, including its used cells.
        result.append({**expected[sheet.title],
                       'grid': {'rowCount': sheet.max_row, 'columnCount': sheet.max_column},
                       'range': f"'{sheet.title.replace(chr(39), chr(39)*2)}'!A1:{get_column_letter(sheet.max_column)}{sheet.max_row}",
                       'values': values})
    workbook.close()
    print(f"Read {book['title']}: {len(result)} registered tabs", flush=True)
    return result


def fetch_all(registry):
    books = registry['spreadsheets']
    if len({b['id'] for b in books}) != len(books):
        raise ValueError('Duplicate source workbook')
    if {t['spreadsheetId'] for t in registry['tabs']} != {b['id'] for b in books}:
        raise ValueError('Source registry mismatch')
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda b: read_book(b, [t for t in registry['tabs'] if t['spreadsheetId'] == b['id']]), books))
    return {'fetchedAt': datetime.now(timezone.utc).isoformat(), 'tabs': [t for batch in results for t in batch]}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    registry = json.loads(Path('tool/sheet_sources.json').read_text(encoding='utf-8-sig'))
    snapshot = fetch_all(registry)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(snapshot, ensure_ascii=False), encoding='utf-8')
    print(f"Complete snapshot: {len(snapshot['tabs'])} tabs")
